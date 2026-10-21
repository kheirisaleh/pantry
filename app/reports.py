import csv
import io
from datetime import timedelta

from . import db, services

PERIODS = {
    "day": "date(purchased_at)",
    "week": "date(purchased_at, 'weekday 0', '-6 days')",
    "month": "strftime('%Y-%m', purchased_at)",
}
MEAL_PERIODS = {
    "day": "date",
    "week": "date(date, 'weekday 0', '-6 days')",
    "month": "strftime('%Y-%m', date)",
}
EVENT_PERIODS = {
    "day": "date(at)",
    "week": "date(at, 'weekday 0', '-6 days')",
    "month": "strftime('%Y-%m', at)",
}


def _range(start, end):
    end = services.parse_date(end) if end else services.today()
    start = services.parse_date(start) if start else end - timedelta(days=89)
    return start.isoformat(), end.isoformat()


def finance(start=None, end=None, period="week"):
    start, end = _range(start, end)
    if period not in PERIODS:
        period = "week"
    window = (start, end + " 23:59:59")
    purchase_filter = "source = 'purchase' AND price_total IS NOT NULL AND purchased_at BETWEEN ? AND ?"
    spending = db.rows(
        "SELECT "
        + PERIODS[period]
        + " AS bucket, ROUND(SUM(price_total), 2) AS total, COUNT(*) AS items FROM batches WHERE "
        + purchase_filter
        + " GROUP BY bucket ORDER BY bucket",
        window,
    )
    stores = db.rows(
        "SELECT COALESCE(store, 'Unknown') AS store, ROUND(SUM(price_total), 2) AS total, COUNT(DISTINCT session_id) AS trips, COUNT(*) AS items FROM batches WHERE "
        + purchase_filter
        + " GROUP BY COALESCE(store, 'Unknown') ORDER BY total DESC",
        window,
    )
    categories = db.rows(
        "SELECT i.category AS category, ROUND(SUM(b.price_total), 2) AS total FROM batches b JOIN ingredients i ON i.id = b.ingredient_id WHERE b.source = 'purchase' AND b.price_total IS NOT NULL AND b.purchased_at BETWEEN ? AND ? GROUP BY i.category ORDER BY total DESC",
        window,
    )
    top_items = db.rows(
        "SELECT i.id, COALESCE(i.name_en, i.raw_name) AS display, i.name_fa, i.icon, ROUND(SUM(b.price_total), 2) AS total, COUNT(*) AS purchases FROM batches b JOIN ingredients i ON i.id = b.ingredient_id WHERE b.source = 'purchase' AND b.price_total IS NOT NULL AND b.purchased_at BETWEEN ? AND ? GROUP BY i.id ORDER BY total DESC LIMIT 15",
        window,
    )
    waste = db.rows(
        "SELECT "
        + EVENT_PERIODS[period]
        + " AS bucket, ROUND(SUM(value), 2) AS total FROM events WHERE kind = 'waste' AND at BETWEEN ? AND ? GROUP BY bucket ORDER BY bucket",
        window,
    )
    waste_items = db.rows(
        "SELECT i.id, COALESCE(i.name_en, i.raw_name) AS display, i.icon, i.unit, ROUND(SUM(e.qty), 1) AS qty, ROUND(SUM(e.value), 2) AS total FROM events e JOIN ingredients i ON i.id = e.ingredient_id WHERE e.kind = 'waste' AND e.at BETWEEN ? AND ? GROUP BY i.id ORDER BY total DESC LIMIT 15",
        window,
    )
    meals = db.rows(
        "SELECT m.id, m.date, m.servings, m.cost, m.kcal, COALESCE(r.name_en, r.raw_name) AS display FROM meals m JOIN recipes r ON r.id = m.recipe_id WHERE m.status = 'cooked' AND m.date BETWEEN ? AND ? ORDER BY m.date DESC",
        (start, end),
    )
    meal_buckets = db.rows(
        "SELECT "
        + MEAL_PERIODS[period]
        + " AS bucket, ROUND(SUM(cost), 2) AS total, COUNT(*) AS meals, ROUND(SUM(cost) / NULLIF(SUM(servings), 0), 2) AS per_serving FROM meals WHERE status = 'cooked' AND date BETWEEN ? AND ? GROUP BY bucket ORDER BY bucket",
        (start, end),
    )
    recipe_costs = db.rows(
        "SELECT r.id, COALESCE(r.name_en, r.raw_name) AS display, COUNT(*) AS times, ROUND(AVG(m.cost / NULLIF(m.servings, 0)), 2) AS per_serving, ROUND(SUM(m.cost), 2) AS total FROM meals m JOIN recipes r ON r.id = m.recipe_id WHERE m.status = 'cooked' AND m.date BETWEEN ? AND ? GROUP BY r.id ORDER BY total DESC",
        (start, end),
    )
    total_spent = round(sum(entry["total"] or 0 for entry in spending), 2)
    total_waste = round(sum(entry["total"] or 0 for entry in waste), 2)
    days = max(1, (services.parse_date(end) - services.parse_date(start)).days + 1)
    return {
        "start": start,
        "end": end,
        "period": period,
        "summary": {
            "spent": total_spent,
            "waste": total_waste,
            "waste_share": round(total_waste / total_spent * 100, 1) if total_spent else 0,
            "per_day": round(total_spent / days, 2),
            "meals_cooked": len(meals),
            "avg_meal_cost": round(sum(meal["cost"] or 0 for meal in meals) / len(meals), 2) if meals else 0,
            "trips": sum(entry["trips"] or 0 for entry in stores),
        },
        "spending": spending,
        "stores": stores,
        "categories": categories,
        "top_items": top_items,
        "waste": waste,
        "waste_items": waste_items,
        "meals": meals[:60],
        "meal_buckets": meal_buckets,
        "recipe_costs": recipe_costs,
    }


def price_trends(limit=40):
    ingredients = db.rows(
        "SELECT i.id, COALESCE(i.name_en, i.raw_name) AS display, i.name_fa, i.icon, i.unit, COUNT(*) AS purchases FROM batches b JOIN ingredients i ON i.id = b.ingredient_id WHERE b.source = 'purchase' AND b.price_total IS NOT NULL AND b.pending_unit IS NULL GROUP BY i.id HAVING purchases >= 1 ORDER BY purchases DESC LIMIT ?",
        (limit,),
    )
    for ingredient in ingredients:
        history = services.price_history(ingredient["id"])
        ingredient["history"] = history
        ingredient["unit_label"] = "per piece" if ingredient["unit"] == "pcs" else ("per kg" if ingredient["unit"] == "g" else "per l")
        if len(history) >= 2:
            first = history[0]["unit_price"]
            last = history[-1]["unit_price"]
            ingredient["change_pct"] = round((last - first) / first * 100, 1) if first else None
        else:
            ingredient["change_pct"] = None
        stores = {}
        for entry in history:
            stores.setdefault(entry["store"] or "Unknown", []).append(entry["unit_price"])
        ingredient["stores"] = [
            {"store": store, "avg": round(sum(values) / len(values), 2), "count": len(values)} for store, values in sorted(stores.items())
        ]
    return ingredients


def nutrition(start=None, end=None):
    start, end = _range(start, end)
    meals = db.rows(
        "SELECT date, SUM(kcal) AS kcal, SUM(protein) AS protein, SUM(fat) AS fat, SUM(carbs) AS carbs, SUM(fiber) AS fiber, COUNT(*) AS meals, SUM(servings) AS servings FROM meals WHERE status = 'cooked' AND date BETWEEN ? AND ? GROUP BY date",
        (start, end),
    )
    snacks = db.rows(
        "SELECT date, SUM(kcal) AS kcal, SUM(protein) AS protein, SUM(fat) AS fat, SUM(carbs) AS carbs, SUM(fiber) AS fiber, COUNT(*) AS snacks FROM snacks WHERE date BETWEEN ? AND ? GROUP BY date",
        (start, end),
    )
    by_day = {}
    for entry in meals:
        by_day[entry["date"]] = {
            "date": entry["date"],
            "meal_kcal": entry["kcal"] or 0,
            "snack_kcal": 0,
            "protein": entry["protein"] or 0,
            "fat": entry["fat"] or 0,
            "carbs": entry["carbs"] or 0,
            "fiber": entry["fiber"] or 0,
            "meals": entry["meals"],
            "servings": entry["servings"] or 0,
        }
    for entry in snacks:
        day = by_day.setdefault(
            entry["date"],
            {
                "date": entry["date"],
                "meal_kcal": 0,
                "snack_kcal": 0,
                "protein": 0,
                "fat": 0,
                "carbs": 0,
                "fiber": 0,
                "meals": 0,
                "servings": 0,
            },
        )
        day["snack_kcal"] += entry["kcal"] or 0
        for key in ("protein", "fat", "carbs", "fiber"):
            day[key] += entry[key] or 0
    days = sorted(by_day.values(), key=lambda entry: entry["date"])
    for day in days:
        day["kcal"] = round(day["meal_kcal"] + day["snack_kcal"])
        for key in ("meal_kcal", "snack_kcal", "protein", "fat", "carbs", "fiber"):
            day[key] = round(day[key], 1)
    weeks = {}
    for day in days:
        monday = (services.parse_date(day["date"]) - timedelta(days=services.parse_date(day["date"]).weekday())).isoformat()
        week = weeks.setdefault(monday, {"week": monday, "kcal": 0, "protein": 0, "fat": 0, "carbs": 0, "fiber": 0, "days": 0})
        week["days"] += 1
        for key in ("kcal", "protein", "fat", "carbs", "fiber"):
            week[key] += day[key]
    week_list = []
    for week in sorted(weeks.values(), key=lambda entry: entry["week"]):
        for key in ("kcal", "protein", "fat", "carbs", "fiber"):
            week[key] = round(week[key], 1)
        week_list.append(week)
    incomplete = db.scalar("SELECT COUNT(*) FROM meals WHERE status = 'cooked' AND date BETWEEN ? AND ? AND kcal IS NULL", (start, end))
    return {"start": start, "end": end, "days": days, "weeks": week_list, "incomplete_meals": incomplete}


def period_report(list_id):
    shopping = db.row("SELECT * FROM shopping_lists WHERE id = ?", (list_id,))
    if not shopping or not shopping["range_start"]:
        raise services.ServiceError("This list has no planning range", 404)
    start = shopping["range_start"]
    end = shopping["range_end"]
    window = (start, end + " 23:59:59")
    bought = db.rows(
        "SELECT b.ingredient_id, SUM(b.qty_initial) AS qty, SUM(b.price_total) AS spent FROM batches b JOIN shopping_sessions s ON s.id = b.session_id WHERE s.list_id = ? AND b.pending_unit IS NULL GROUP BY b.ingredient_id",
        (list_id,),
    )
    ids = {entry["ingredient_id"] for entry in bought}
    used = {
        entry["ingredient_id"]: entry["qty"]
        for entry in db.rows(
            "SELECT ingredient_id, SUM(qty) AS qty FROM events WHERE kind IN ('cook', 'snack') AND at BETWEEN ? AND ? GROUP BY ingredient_id",
            window,
        )
    }
    wasted = {
        entry["ingredient_id"]: entry["qty"]
        for entry in db.rows(
            "SELECT ingredient_id, SUM(qty) AS qty FROM events WHERE kind = 'waste' AND at BETWEEN ? AND ? GROUP BY ingredient_id", window
        )
    }
    short = {
        entry["ingredient_id"]: entry["qty"]
        for entry in db.rows(
            "SELECT ingredient_id, SUM(qty) AS qty FROM events WHERE kind IN ('cook_shortfall', 'snack_shortfall') AND at BETWEEN ? AND ? GROUP BY ingredient_id",
            window,
        )
    }
    ids |= set(short)
    stock = services.stock_map()
    ingredients = services.ingredient_map(ids)
    lines = []
    for ingredient_id in ids:
        ingredient = ingredients.get(ingredient_id)
        if not ingredient:
            continue
        entry = next((item for item in bought if item["ingredient_id"] == ingredient_id), {"qty": 0, "spent": 0})
        lines.append(
            {
                "ingredient_id": ingredient_id,
                "display": services.display_name(ingredient),
                "icon": ingredient.get("icon"),
                "unit": ingredient["unit"],
                "bought": entry["qty"] or 0,
                "spent": round(entry["spent"] or 0, 2),
                "used": used.get(ingredient_id, 0),
                "wasted": wasted.get(ingredient_id, 0),
                "short": short.get(ingredient_id, 0),
                "left": stock.get(ingredient_id, 0),
            }
        )
    lines.sort(key=lambda line: (-line["short"], -line["left"], line["display"]))
    planned = db.scalar("SELECT COUNT(*) FROM meals WHERE date BETWEEN ? AND ?", (start, end))
    cooked = db.scalar("SELECT COUNT(*) FROM meals WHERE status = 'cooked' AND date BETWEEN ? AND ?", (start, end))
    return {
        "list_id": list_id,
        "start": start,
        "end": end,
        "planned_meals": planned,
        "cooked_meals": cooked,
        "shortages": [line for line in lines if line["short"] > 0],
        "surplus": [line for line in lines if line["left"] > 0 and line["short"] <= 0],
        "lines": lines,
    }


def periods():
    return db.rows(
        "SELECT id, range_start, range_end, status, closed_at FROM shopping_lists WHERE range_start IS NOT NULL ORDER BY id DESC LIMIT 30"
    )


EXPORTS = {
    "purchases": (
        ["date", "store", "item", "category", "quantity", "unit", "price_eur"],
        "SELECT b.purchased_at, b.store, COALESCE(i.name_en, i.raw_name), i.category, ROUND(b.qty_initial, 2), COALESCE(b.pending_unit, i.unit), b.price_total FROM batches b JOIN ingredients i ON i.id = b.ingredient_id WHERE b.source = 'purchase' ORDER BY b.purchased_at",
    ),
    "meals": (
        ["date", "dish", "servings", "status", "cook", "cost_eur", "kcal", "protein_g", "fat_g", "carbs_g", "fiber_g"],
        "SELECT m.date, COALESCE(r.name_en, r.raw_name), m.servings, m.status, u.name, m.cost, m.kcal, m.protein, m.fat, m.carbs, m.fiber FROM meals m JOIN recipes r ON r.id = m.recipe_id LEFT JOIN users u ON u.id = m.cooked_by ORDER BY m.date, m.position",
    ),
    "snacks": (
        ["date", "item", "quantity", "unit", "user", "kcal", "protein_g", "cost_eur"],
        "SELECT s.date, COALESCE(i.name_en, i.raw_name), ROUND(s.qty, 2), i.unit, u.name, s.kcal, s.protein, s.cost FROM snacks s JOIN ingredients i ON i.id = s.ingredient_id LEFT JOIN users u ON u.id = s.user_id ORDER BY s.date",
    ),
    "waste": (
        ["date", "item", "quantity", "unit", "value_eur", "user"],
        "SELECT e.at, COALESCE(i.name_en, i.raw_name), ROUND(e.qty, 2), i.unit, ROUND(e.value, 2), u.name FROM events e JOIN ingredients i ON i.id = e.ingredient_id LEFT JOIN users u ON u.id = e.user_id WHERE e.kind = 'waste' ORDER BY e.at",
    ),
    "inventory": (
        ["item", "persian", "category", "kind", "quantity", "unit"],
        "SELECT COALESCE(i.name_en, i.raw_name), i.name_fa, i.category, i.kind, ROUND(COALESCE(SUM(b.qty_remaining), 0), 2), i.unit FROM ingredients i LEFT JOIN batches b ON b.ingredient_id = i.id AND b.pending_unit IS NULL WHERE i.archived = 0 GROUP BY i.id ORDER BY i.category, 1",
    ),
    "recipes": (
        ["dish", "persian", "ingredient", "amount_per_person", "unit", "optional", "default_choice"],
        "SELECT COALESCE(r.name_en, r.raw_name), r.name_fa, COALESCE(i.name_en, i.raw_name), ri.amount, ri.unit, ri.optional, ri.is_default FROM recipe_items ri JOIN recipes r ON r.id = ri.recipe_id JOIN ingredients i ON i.id = ri.ingredient_id WHERE r.archived = 0 ORDER BY 1, ri.slot, ri.position",
    ),
}


def export_csv(name):
    if name not in EXPORTS:
        raise services.ServiceError("Unknown export", 404)
    header, sql = EXPORTS[name]
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    for record in db.conn().execute(sql).fetchall():
        writer.writerow(list(record))
    return buffer.getvalue()


def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (header, sql) in EXPORTS.items():
        sheet = workbook.create_sheet(name.capitalize())
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2430")
        for record in db.conn().execute(sql).fetchall():
            sheet.append(list(record))
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(10, width + 2))
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

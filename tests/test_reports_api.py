import csv
import io

from openpyxl import load_workbook

from app import db


def shop_and_cook(client, kitchen):
    recipe = kitchen.recipe()
    kitchen.meal(recipe["id"], 0, servings=2)
    built = client.post("/api/shopping/build", json={"start": kitchen.day(0), "end": kitchen.day(1)}).json()["list"]
    session = client.post("/api/shopping/sessions", json={"store": "Rewe"}).json()
    prices = {"Basmati rice": 2.0, "Ground beef": 4.0, "Green beans": 1.0}
    for item in built["items"]:
        client.patch("/api/shopping/items/" + str(item["id"]), json={"checked": True, "price": prices[item["ingredient"]["display"]]})
    client.post("/api/shopping/sessions/" + str(session["session_id"]) + "/finish")
    client.post("/api/inventory/" + str(kitchen.ingredient_id("Green beans")) + "/waste", json={"qty": 10, "unit": "g"})
    meal_id = db.scalar("SELECT id FROM meals")
    client.post("/api/cook", json={"meal_id": meal_id})
    return built


def test_finance(client, kitchen):
    shop_and_cook(client, kitchen)
    report = client.get("/api/reports/finance", params={"period": "day"}).json()
    summary = report["summary"]
    assert summary["spent"] == 7.0 and summary["meals_cooked"] == 1 and summary["trips"] == 1
    assert summary["waste"] == round(10 / 154, 2)
    assert report["stores"][0]["store"] == "Rewe" and report["stores"][0]["total"] == 7.0
    assert {entry["category"] for entry in report["categories"]} == {"grains", "meat", "vegetables"}
    assert report["top_items"][0]["display"] == "Ground beef"
    assert report["spending"][0]["bucket"] == kitchen.day(0)
    assert report["recipe_costs"][0]["display"] == "Loobia Polo"
    assert client.get("/api/reports/finance", params={"period": "decade"}).json()["period"] == "week"
    month = client.get("/api/reports/finance", params={"period": "month"}).json()
    assert month["spending"][0]["bucket"] == kitchen.day(0)[:7]
    assert client.get("/api/reports/finance", params={"start": "bad"}).status_code == 400


def test_prices_and_nutrition(client, kitchen):
    shop_and_cook(client, kitchen)
    prices = client.get("/api/reports/prices").json()
    rice = next(entry for entry in prices if entry["display"] == "Basmati rice")
    assert rice["unit_label"] == "per kg" and rice["change_pct"] is None
    assert rice["history"][0]["unit_price"] == 10.0 and rice["stores"] == [{"store": "Rewe", "avg": 10.0, "count": 1}]
    nutrition = client.get("/api/reports/nutrition").json()
    assert nutrition["days"][0]["meals"] == 1 and nutrition["incomplete_meals"] == 0
    assert nutrition["days"][0]["kcal"] == round(2 * (350 + 200 + 70 * 0.31))
    assert nutrition["weeks"][0]["days"] == 1


def test_periods(client, kitchen):
    built = shop_and_cook(client, kitchen)
    periods = client.get("/api/reports/periods").json()
    assert periods[0]["id"] == built["id"] and periods[0]["status"] == "open"
    report = client.get("/api/reports/period/" + str(built["id"])).json()
    assert report["planned_meals"] == 1 and report["cooked_meals"] == 1
    lines = {line["display"]: line for line in report["lines"]}
    assert lines["Basmati rice"]["bought"] == 200.0 and lines["Basmati rice"]["used"] == 200.0
    assert lines["Green beans"]["wasted"] == 10.0
    assert client.get("/api/reports/period/99").status_code == 404


def test_csv_and_xlsx(client, kitchen):
    shop_and_cook(client, kitchen)
    for name in ("purchases", "meals", "snacks", "waste", "inventory", "recipes"):
        response = client.get("/api/export/" + name + ".csv")
        assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
        assert "attachment" in response.headers["content-disposition"]
        rows = list(csv.reader(io.StringIO(response.text)))
        assert rows[0]
    purchases = list(csv.reader(io.StringIO(client.get("/api/export/purchases.csv").text)))
    assert purchases[0] == ["date", "store", "item", "category", "quantity", "unit", "price_eur"] and len(purchases) == 4
    assert client.get("/api/export/secrets.csv").status_code == 404
    workbook = load_workbook(io.BytesIO(client.get("/api/export.xlsx").content))
    assert workbook.sheetnames == ["Purchases", "Meals", "Snacks", "Waste", "Inventory", "Recipes"]
    assert workbook["Meals"]["B2"].value == "Loobia Polo"
